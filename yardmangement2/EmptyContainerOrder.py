#! python3

import os
from tkinter import filedialog as FileDialog
from tkinter import messagebox as MessageBox
from tkinter import *  # Prueba borrar cuadro dialogo
from tkinter.ttk import * # Prueba borrar cuadro dialogo
from os.path import basename
import openpyxl
import xlrd
import xlwt
from xlutils.copy import copy
import re
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd
import pyexcel as p
from xls2xlsx import XLS2XLSX
from operator import itemgetter, attrgetter
from datetime import datetime


pathFile = r"S:/OPS/PLANIFICACIÓN/Secuencia de retiro vacíos/"
#print(pathFile)
root = Tk() # ERASE DIALOG
root.update()


def clear():  # ERASE SCREEN
    if os.name == "nt":      # Windows
        os.system("cls")
    else:
        os.system("clear")   # Linux


def openFile():  # DIALOG TO CHOOSE A FILE
    excel_file = FileDialog.askopenfilename(
        initialdir=pathFile,
        filetypes=(("xls", "*.xls"), ("all files", "*.*")),
        title="Escoja el archivo de Existencias Excel a Procesar"
    )
    base = os.path.basename(excel_file)
    name = os.path.splitext(base)[0]
    print("\nEl nombre del archivo inicial es: " + name + "\n")
    return(excel_file)


def askUser(): # DIALOG TO CHOOSE VESSEL
   buque = str((input("Digite el buque del cual generar el listado: ").upper()))
   #print(buque)
   return(buque)


def ordered_list(sheet, buque_seleccion, option): # RECEIVE EXCEL SHIFT AND GENERETE A ORDERED LIST

    #INIT VAR AND LIST
    mty_containers_ABC_even = []
    mty_containers_ABC_odd = []
    mty_containers_Z = []
    now = datetime.now()
    format = now.strftime("%d-%m-%Y %H-%M")

    # Loop through the sheet
    for i in range(sheet.nrows):
        ubicacion = (repr(sheet.cell_value(i, 21)).replace("'",""))
        contenedor = (repr(sheet.cell_value(i, 0)).replace("'",""))
        pies = (repr(sheet.cell_value(i, 3)))
        tipo = (repr(sheet.cell_value(i, 2)).replace("'",""))
        estatus = (repr(sheet.cell_value(i, 9)).replace("'",""))
        buque_import = (repr(sheet.cell_value(i, 11)).replace("'",""))
        viaje_import = (repr(sheet.cell_value(i, 12)).replace("'",""))
        buque_export = (repr(sheet.cell_value(i, 13)).replace("'",""))
        viaje_export = (repr(sheet.cell_value(i, 14)).replace("'",""))
        zona = (ubicacion[0:1])
        bloque = (ubicacion[2:4])
        modulo = (ubicacion[5:8])
        calle = (ubicacion[9:12])
        altura = (ubicacion[12:13])
        
        if option == 1: # Buque import       
            if buque_import == buque_seleccion:
                file_ordered = ("secuencia retiro " + buque_seleccion + " " + format + ".xlsx") # name of final file
                if estatus == "EMT" or estatus == "TRV":
                    line = {"contenedor" : contenedor, "pies" : pies, "tipo": tipo, "estatus" : estatus,
                               "buque_import" : buque_import, "viaje_import" : viaje_import, "ubicacion" : ubicacion,
                               "zona": zona, "bloque" : bloque, "modulo" : modulo, "calle" : calle, "altura" : altura}
                    if zona == "A" or zona == "B" or zona =="C":
                        if int(bloque) % 2 == 0:
                            mty_containers_ABC_even.append(line)
                        elif int(bloque) % 2 == 1:
                            mty_containers_ABC_odd.append(line)
                    elif zona == "Z":
                        mty_containers_Z.append(line)
        elif option == 2: # Buque export
               if buque_export == buque_seleccion:
                    file_ordered = ("secuencia retiro " + " " + format + ".xlsx") # name of final file
                    if estatus == "EMT" or estatus == "TRV":
                        line = {"contenedor" : contenedor, "pies" : pies, "tipo": tipo, "estatus" : estatus,
                                   "ubicacion" : ubicacion, "zona": zona, "bloque" : bloque, "modulo" : modulo, "calle" : calle, "altura" : altura}
                        if zona == "A" or zona == "B" or zona =="C":
                            if int(bloque) % 2 == 0:
                                mty_containers_ABC_even.append(line)
                            elif int(bloque) % 2 == 1:
                                mty_containers_ABC_odd.append(line)
                        elif zona == "Z":
                            mty_containers_Z.append(line)
                

    # Order lists
    if len(mty_containers_ABC_even) > 0:  #order list par           
        #mty_containers_ABC_even.sort(key=lambda x: x.get("modulo"), reverse=True)
        mty_containers_ABC_even.sort(key=lambda x: x.get("calle"))
        mty_containers_ABC_even.sort(key=lambda x: x.get("altura"), reverse=True)
        mty_containers_ABC_even.sort(key=lambda x: x.get("modulo"))
        mty_containers_ABC_even.sort(key=lambda x: x.get("bloque"))
       
    if len(mty_containers_ABC_odd) > 0: #order list odd           
        #mty_containers_ABC_even.sort(key=lambda x: x.get("modulo"), reverse=True)
        mty_containers_ABC_odd.sort(key=lambda x: x.get("calle"), reverse=True)
        mty_containers_ABC_odd.sort(key=lambda x: x.get("altura"), reverse=True)
        mty_containers_ABC_odd.sort(key=lambda x: x.get("modulo"))
        mty_containers_ABC_odd.sort(key=lambda x: x.get("bloque"))
        
    if len(mty_containers_Z) > 0: # order list Z        
        mty_containers_Z.sort(key=lambda x: x.get("altura"), reverse=True)
        mty_containers_Z.sort(key=lambda x: x.get("calle"), reverse=True)
        mty_containers_Z.sort(key=lambda x: x.get("modulo"), reverse=True)
        mty_containers_Z.sort(key=lambda x: x.get("bloque"), reverse=True)


    # Unir todas las listas ordenadas
    mty_containers =  mty_containers_ABC_odd + mty_containers_ABC_even + mty_containers_Z

    if len(mty_containers) > 0:
        print (f"\nTotal unidades en el listado: {len(mty_containers)}")
        if option == 1:
            # Generar archivo
            df = pd.DataFrame(mty_containers, columns = ["contenedor", "pies", "tipo", "estatus",
                                "buque_import", "viaje_import", "ubicacion", "zona", "bloque",
                                "modulo", "calle", "altura"])
            df.index.name = "Secuencia"
            df.index = df.index + 1 # index begin with 1 and not 0
        elif option == 2:
            # Generar archivo
            df = pd.DataFrame(mty_containers, columns = ["contenedor", "pies", "tipo", "estatus",
                                "ubicacion", "zona", "bloque", "modulo", "calle", "altura"])
            df.index.name = "Secuencia"
            df.index = df.index + 1 # index begin with 1 and not 0
            
             
        os.chdir(pathFile) # go to Pathhfile
        #base = os.path.basename(excel_file)
        #name = os.path.splitext(base)[0]
        #df.to_csv(name + ".csv", sep=";") # Save like CSV
        #file_ordered = ("secuencia retiro " + buque_seleccion + " " + format + ".xlsx")
        #df.to_excel(file_ordered) # Crea directamente el libro excel

        # FORMAT CELLS
        writer = pd.ExcelWriter(file_ordered, engine='xlsxwriter')
        df.to_excel(writer, sheet_name="Secuencia_Retiro")
        workbook = writer.book
        worksheet = writer.sheets["Secuencia_Retiro"]

        if option == 1:
            # format-font
            cell_format = workbook.add_format()
            cell_format.set_font_name('Verdana')
            cell_format.set_font_size(11)
            worksheet.set_column("A:M", None, cell_format)
            # Columns     
            worksheet.set_column("B:B", 21)
            worksheet.set_column("C:C", 7)
            worksheet.set_column("D:D", 7)
            worksheet.set_column("E:E", 9)
            worksheet.set_column("F:F", 15)
            worksheet.set_column("G:G", 18)
            worksheet.set_column("H:H", 18)
            worksheet.set_column("I:M", 9)
        elif option == 2:
            # format-font
            cell_format = workbook.add_format()
            cell_format.set_font_name('Verdana')
            cell_format.set_font_size(11)
            worksheet.set_column("A:K", None, cell_format)
            # Columns     
            worksheet.set_column("B:B", 21)
            worksheet.set_column("C:C", 7)
            worksheet.set_column("D:D", 7)
            worksheet.set_column("E:E", 9)
            worksheet.set_column("F:F", 18)
            worksheet.set_column("G:K", 9)    
                  
        (max_row, max_col) = df.shape
        # Create a list of column headers, to use in add_table().
        column_settings = [{'header': column} for column in df.columns]
        # Add the Excel table structure. Pandas will add the data.
        worksheet.add_table(0, 1, max_row, max_col, {'columns': column_settings})
        # Make the columns wider for clarity.
        worksheet.set_column(0, max_col + 1, 12)  

        writer.save()
              
    else:
        print("\nEl buque seleccionado no coincide con la búsqueda o no tiene unidades vacías para asignar una secuencia de retiro")
        ornament()
        

def ornament():
    print("_" * 50)


def menu1():
    ornament()
    print("Programa para ordenar el retiro de vacíos por alturas y/ó filas")
    ornament()
      
    
def menu2():
    #print("\n")
    ornament()
    print("Opciones disponibles:")
    ornament()
    print("1. Vacíos descargados desde Motonaves")
    print("2. Vacíos desde el FreePool")
    print("0. Salir")
    print()

  
def main():  # MAIN
    clear()  
    menu1()
    excel_file= openFile()  # Name of the file to choose
    root.destroy() # Close Dialog
    buque_seleccion =  askUser()
    wb = xlrd.open_workbook(excel_file, encoding_override="cp1252") # Read excel book
    sheet = wb.sheet_by_index(0)

    while True:       
            option = menu2()
            try:
                option =  int(input("Seleccione una opción: "))
                ornament()
                if option == 1: # Vacíos desde buque
                    ordered_list(sheet, buque_seleccion, option)
                    break
                elif option == 2:
                    ordered_list(sheet, buque_seleccion, option)
                    break
                elif option == 0:
                    break
                else:
                    print()
                    clear()
                    print("Error, solo de aceptan números del 0 al 2")
            except ValueError:
                  print("Error, ingrese solamente números")            
    

    


if __name__ == '__main__':
    main()


os.system("pause") # Press a key to continue 
