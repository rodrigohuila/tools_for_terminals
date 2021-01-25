#! python3

import os
from tkinter import filedialog as FileDialog
from tkinter import messagebox as MessageBox
from tkinter import *  # Prueba borrar cuadro dialogo
from tkinter.ttk import * # Prueba borrar cuadro dialogo
from os.path import basename
import openpyxl
import xlrd
import xlwt
from xlutils.copy import copy
import re
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd
import pyexcel as p
from xls2xlsx import XLS2XLSX
from operator import itemgetter, attrgetter
from datetime import datetime, date, timedelta
import time
from ocupacion import ocupacion_list, count_teus
from quantities_by_type import quantities_list, count_quantity
from CopyPreestibaFromExcel import getPreestibas, greetings, writePreestibas, getFile
import os.path
import pyautogui



pathFile = r"S:/OPS/PLANIFICACION/Listados Yard Planning/"
#print(pathFile)
#root = Tk() # ERASE DIALOG
#root.update()
espacios = 80


def clear():  # ERASE SCREEN
    if os.name == "nt":      # Windows
        os.system("cls")
    else:
        os.system("clear")   # Linux


def openFile():  # DIALOG TO CHOOSE A FILE
    excel_file = FileDialog.askopenfilename(
        initialdir=pathFile,
        filetypes=(("xls", "*.xls"), ("all files", "*.*")),
        title="Escoja el archivo de Existencias Excel a Procesar"
    )
    base = os.path.basename(excel_file)
    name = os.path.splitext(base)[0]
    print("\nEl nombre del archivo inicial es: " + name + "\n")
    return(excel_file)


def askUser(option): # DIALOG TO CHOOSE VESSEL
    if option == 1 or option == 6:
        seleccion = str((input("Digite el buque del cual generar el listado: ").upper()))       
    if option == 2:
        seleccion = str((input("Digite el Dummie del cual generar el listado: ").upper()))          
    elif option == 3:
        seleccion = str((input("Digite la Línea del cual generar el listado: ").upper()))        
    elif option == 4:
        seleccion = str((input("Digite el puerto de descarga del cual generar el listado: ").upper()))        
    elif option == 5:
        seleccion = str((input("Digite el número de la retención del cual generar el listado: ").upper()))
    return(seleccion)

def typeContainer(file):
    preestibas = getPreestibas(file)
    book_Pres = xlrd.open_workbook(file)
    sheet_Pres = book_Pres.sheet_by_index(1)
   
    greetings()

    opc = input ('\nEscoja Monitor No 1 o Monitor No 2: ')

    ornament()
    print ('Cant   Contenedor')
    ornament()

    for i in range(sheet_Pres.nrows):  # Loop through the sheet
        contenedor = (repr(sheet_Pres.cell_value(i, 0)).replace("'", ""))
        if type(preestibas) == list:
            if contenedor in preestibas:
                writePreestibas(int(opc), contenedor, i)
                
    print('\nTotal preestibas digitadas:' + str(len(preestibas)))           

    if int(opc) == 1:
        pyautogui.click(1168, 606) # clic Cancelar Monitor 1
    elif int(opc) == 2 :
        pyautogui.click(3095, 606) # clic Cancelar Monitor 2
        

def ordered_list(sheet, seleccion, option): # RECEIVE EXCEL SHIFT AND GENERETE A ORDERED LIST

    #INIT VAR AND LIST
    mty_containers_ABC_even = []
    mty_containers_ABC_odd = []
    mty_containers_Z = []
    mtycontainers_linea = []
    mty_asignacion = []
    now = datetime.now()
    format = now.strftime("%d-%m-%Y %H-%M")
           
    # Loop through the sheet
    for i in range(sheet.nrows):
        ubicacion = (repr(sheet.cell_value(i, 21)).replace("'",""))
        contenedor = (repr(sheet.cell_value(i, 0)).replace("'",""))
        pies = (repr(sheet.cell_value(i, 3)))
        tipo = (repr(sheet.cell_value(i, 2)).replace("'",""))
        estatus = (repr(sheet.cell_value(i, 9)).replace("'",""))
        buque_import = (repr(sheet.cell_value(i, 11)).replace("'",""))
        viaje_import = (repr(sheet.cell_value(i, 12)).replace("'",""))
        buque_export = (repr(sheet.cell_value(i, 13)).replace("'",""))
        viaje_export = (repr(sheet.cell_value(i, 14)).replace("'",""))
        zona = (ubicacion[0:1])
        bloque = (ubicacion[2:4])
        modulo = (ubicacion[5:8])
        calle = (ubicacion[9:12])
        altura = (ubicacion[12:13])
        linea = (repr(sheet.cell_value(i, 10)).replace("'",""))
        iso = (repr(sheet.cell_value(i, 1)).replace("'",""))
        dias_terminal = repr(sheet.cell_value(i, 24)).replace("'","")
        pdescarga = (repr(sheet.cell_value(i, 6)).replace("'",""))
        pfinal = (repr(sheet.cell_value(i, 7)).replace("'",""))
        retencion = (repr(sheet.cell_value(i, 19)).replace("'",""))
        observacion = (repr(sheet.cell_value(i, 26)).replace("'",""))
        sit = (repr(sheet.cell_value(i, 8)).replace("'",""))      

        
        if option == 1: # Buque import       
            if buque_import == seleccion:
                file_ordered = ("secuencia retiro " + seleccion + " " + format + ".xlsx") # name of final file
                if estatus == "EMT" or estatus == "TRV":
                    pies = int(float(pies))
                    line = {"contenedor" : contenedor, "pies" : pies, "tipo": tipo, "estatus" : estatus,
                               "buque_import" : buque_import, "viaje_import" : viaje_import, "ubicacion" : ubicacion,
                               "zona": zona, "bloque" : bloque, "modulo" : modulo, "calle" : calle, "altura" : altura}
                    if zona == "A" or zona == "B" or zona =="C":
                        if int(bloque) % 2 == 0:
                            mty_containers_ABC_even.append(line)
                        elif int(bloque) % 2 == 1:
                            mty_containers_ABC_odd.append(line)
                    elif zona == "Z":
                        mty_containers_Z.append(line)
        elif option == 2: # Buque export
               if buque_export == seleccion:
                    file_ordered = ("secuencia retiro " + seleccion + " " + format + ".xlsx") # name of final file
                    if estatus == "EMT" or estatus == "TRV":
                        pies = int(float(pies))
                        line = {"contenedor" : contenedor, "pies" : pies, "tipo": tipo, "estatus" : estatus,
                                   "ubicacion" : ubicacion, "zona": zona, "bloque" : bloque, "modulo" : modulo, "calle" : calle, "altura" : altura}
                        if zona == "A" or zona == "B" or zona =="C":
                            if int(bloque) % 2 == 0:
                                mty_containers_ABC_even.append(line)
                            elif int(bloque) % 2 == 1:
                                mty_containers_ABC_odd.append(line)
                        elif zona == "Z":
                            mty_containers_Z.append(line)
        elif option == 3: # By Line
              if linea == seleccion:
                    file_ordered = ("Unidades en Existencia" + " " + linea + " " + format + ".xlsx") # name of final file
                    if estatus == "EMT" or estatus == "TRV":
                        pies = int(float(pies))
                        dias_terminal = int(float(dias_terminal))
                        #print(dias_terminal)
                        fecha = ((now - timedelta(days=dias_terminal)).strftime("%d-%m-%Y"))
                        line = {"contenedor" : contenedor, "pies" : pies, "tipo": tipo, "iso" : iso, "linea" : linea, "estatus" : estatus,
                                   "ubicacion" : ubicacion, "fecha" : fecha}
                        mtycontainers_linea.append(line)
        elif option == 4: # By Port
            if pdescarga == seleccion:
                file_ordered = (f"Listado {pdescarga} " + format + ".xlsx") # name of final file
                if retencion =="":
                    if estatus == "EMT":
                        pies = int(float(pies))
                        dias_terminal = int(float(dias_terminal))
                        fecha = ((now - timedelta(days=dias_terminal)).strftime("%d-%m-%Y"))
                        line = {"contenedor" : contenedor, "pies" : pies, "tipo": tipo, "pdescarga": pdescarga, "linea" : linea, "estatus" : estatus,
                                   "ubicacion" : ubicacion, "dias_en_terminal" : dias_terminal, "fecha_ingreso": fecha, "zona": zona, "bloque" : bloque,
                                   "modulo" : modulo, "calle" : calle, "altura" : altura}
                        if zona == "A" or zona == "B" or zona =="C":
                            if int(bloque) % 2 == 0:
                                mty_containers_ABC_even.append(line)
                            elif int(bloque) % 2 == 1:
                                mty_containers_ABC_odd.append(line)
                        elif zona == "Z":
                            mty_containers_Z.append(line)
        elif option == 5: # By Retención
            if retencion == seleccion:
                file_ordered = (f"Listado unidades retención {retencion} " + format + ".xlsx") # name of final file
                if retencion == seleccion:
                    if estatus == "EMT":
                        pies = int(float(pies))
                        dias_terminal = int(float(dias_terminal))
                        fecha = ((now - timedelta(days=dias_terminal)).strftime("%d-%m-%Y"))
                        line = {"contenedor" : contenedor, "pies" : pies, "tipo": tipo, "pdescarga": pdescarga, "linea" : linea, "estatus" : estatus,
                                   "ubicacion" : ubicacion, "observacion": observacion, "dias_en_terminal" : dias_terminal, "fecha_ingreso": fecha, "zona": zona, "bloque" : bloque,
                                   "modulo" : modulo, "calle" : calle, "altura" : altura}
                        if zona == "A" or zona == "B" or zona =="C" or zona =="F":
                            if int(bloque) % 2 == 0:
                                mty_containers_ABC_even.append(line)
                            elif int(bloque) % 2 == 1:
                                mty_containers_ABC_odd.append(line)
                        elif zona == "Z" or zona == "H" or zona == "P" or zona == "S" or zona == "M":
                            mty_containers_Z.append(line)
        elif option == 6: # Asignación Evacuación
            if buque_export == seleccion:
                file_ordered = (f"Vacíos asignados {seleccion} {format}.xlsx") # name of final file
                if estatus == "EMT" or estatus == "TRV":
                    if sit == "C":
                        pies = int(float(pies))
                        quantity = count_quantity(pies)
                        dias_terminal = int(float(dias_terminal))
                        fecha = ((now - timedelta(days=dias_terminal)).strftime("%d-%m-%Y"))
                        line = {"contenedor": contenedor, "pies": pies, "tipo": tipo, "estatus": estatus,"pdescarga": pdescarga, "pfinal":  pfinal,
                                   "sit": sit, "linea":linea, "buque_export" : buque_export, "viaje_export" : viaje_export , "fecha_ingreso": fecha, "conts": quantity}
                        mty_asignacion.append(line)
                        viaje = viaje_export
                        
                                          
    # Order lists
    if len(mty_containers_ABC_even) > 0:  #order list par           
        mty_containers_ABC_even.sort(key=lambda x: x.get("calle")) # 1 -> 6
        mty_containers_ABC_even.sort(key=lambda x: x.get("altura"), reverse=True) # F -> A
        mty_containers_ABC_even.sort(key=lambda x: x.get("modulo"))
        mty_containers_ABC_even.sort(key=lambda x: x.get("bloque"))       
    if len(mty_containers_ABC_odd) > 0: #order list odd           
        mty_containers_ABC_odd.sort(key=lambda x: x.get("calle"), reverse=True) # 6 -> 1
        mty_containers_ABC_odd.sort(key=lambda x: x.get("altura"), reverse=True) # F -> A
        mty_containers_ABC_odd.sort(key=lambda x: x.get("modulo"))
        mty_containers_ABC_odd.sort(key=lambda x: x.get("bloque"))        
    if len(mty_containers_Z) > 0: # order list Z        
        mty_containers_Z.sort(key=lambda x: x.get("altura"), reverse=True)
        mty_containers_Z.sort(key=lambda x: x.get("calle"), reverse=True)
        mty_containers_Z.sort(key=lambda x: x.get("modulo"), reverse=True)
        mty_containers_Z.sort(key=lambda x: x.get("bloque"), reverse=True)
    if len(mtycontainers_linea) > 0: # order list líneas        
        mtycontainers_linea.sort(key=lambda x: x.get("fecha"), reverse=False)
        mtycontainers_linea.sort(key=lambda x: x.get("contenedor"), reverse=False)
        mtycontainers_linea.sort(key=lambda x: x.get("tipo"), reverse=False)
        mtycontainers_linea.sort(key=lambda x: x.get("pies"), reverse=False)
    if len(mty_asignacion) > 0: # order asignación:
        mty_asignacion.sort(key=lambda x: x.get("contenedor"), reverse=False)
        mty_asignacion.sort(key=lambda x: x.get("tipo"), reverse=False)
        mty_asignacion.sort(key=lambda x: x.get("pies"), reverse=False)
        mty_asignacion.sort(key=lambda x: x.get("pdescarga"), reverse=False)
        dfasignacion = pd.DataFrame(mty_asignacion)    
        

    # Unir todas las listas ordenadas
    mty_containers =  mty_containers_ABC_odd + mty_containers_ABC_even + mty_containers_Z + mtycontainers_linea + mty_asignacion

    if len(mty_containers) > 0:
        print (f"\nTotal unidades en el listado: {len(mty_containers)}")
        if option == 1:
            # Generar archivo
            df = pd.DataFrame(mty_containers, columns = ["contenedor", "pies", "tipo", "estatus",
                                "buque_import", "viaje_import", "ubicacion", "zona", "bloque",
                                "modulo", "calle", "altura"])
            df.index.name = "Secuencia"
            df.index = df.index + 1 # index begin with 1 and not 0
        elif option == 2:
            # Generar archivo
            df = pd.DataFrame(mty_containers, columns = ["contenedor", "pies", "tipo", "estatus",
                                "ubicacion", "zona", "bloque", "modulo", "calle", "altura"])
            df.index.name = "Secuencia"
            df.index = df.index + 1 # index begin with 1 and not 0
        elif option == 3:
            # Generar archivo
            df = pd.DataFrame(mty_containers, columns = ["contenedor", "pies", "tipo", "iso", "linea",
                                "estatus", "fecha"])
            df.index.name = "cant"
            df.index = df.index + 1 # index begin with 1 and not 0
        elif option == 4:
            # Generar archivo
            df = pd.DataFrame(mty_containers, columns = ["contenedor", "pies", "tipo", "pdescarga", "linea", "estatus",
                                "ubicacion", "dias_en_terminal", "fecha_ingreso"])
            df.index.name = "cant"
            df.index = df.index + 1 # index begin with 1 and not 0
        elif option == 5:
            # Generar archivo
            df = pd.DataFrame(mty_containers, columns = ["contenedor", "pies", "tipo", "pdescarga", "linea", "estatus",
                                "ubicacion", "observacion", "dias_en_terminal", "fecha_ingreso"])
            df.index.name = "cant"
            df.index = df.index + 1 # index begin with 1 and not 0
        elif option == 6:
             # Generar archivo
             df = pd.DataFrame(mty_containers, columns = ["contenedor", "pies", "tipo", "estatus",
                                 "pdescarga", "pfinal", "sit", "linea", "buque_export", "viaje_export", "fecha_ingreso"])
             df.index.name = "cant"
             df.index = df.index + 1 # index begin with 1 and not 0    

            
        os.chdir(pathFile) # go to Pathhfile
        #base = os.path.basename(excel_file)
        #name = os.path.splitext(base)[0]
        #df.to_csv(name + ".csv", sep=";") # Save like CSV
        #df.to_excel(file_ordered) # Crea directamente el libro excel

        # FORMAT CELLS
        writer = pd.ExcelWriter(file_ordered, engine='xlsxwriter')
        df.to_excel(writer, sheet_name=seleccion)
        workbook = writer.book
        worksheet = writer.sheets[seleccion]

        if option == 1:
            # format-font
            cell_format = workbook.add_format()
            cell_format.set_font_name('Verdana')
            cell_format.set_font_size(11)
            worksheet.set_column("A:M", None, cell_format)
            # Columns
            worksheet.set_column("A:A", 9.5)
            worksheet.set_column("B:B", 18.8)
            worksheet.set_column("C:C", 7)
            worksheet.set_column("D:D", 7)
            worksheet.set_column("E:E", 9)
            worksheet.set_column("F:F", 15)
            worksheet.set_column("G:G", 18)
            worksheet.set_column("H:H", 18)
            worksheet.set_column("I:M", 9)
        elif option == 2:
            # format-font
            cell_format = workbook.add_format()
            cell_format.set_font_name('Verdana')
            cell_format.set_font_size(11)
            worksheet.set_column("A:K", None, cell_format)
            # Columns
            worksheet.set_column("A:A", 9.5)
            worksheet.set_column("B:B", 18.8)
            worksheet.set_column("C:C", 7)
            worksheet.set_column("D:D", 7)
            worksheet.set_column("E:E", 9)
            worksheet.set_column("F:F", 18)
            worksheet.set_column("G:K", 9)
        elif option == 3:
            # format-font
            cell_format = workbook.add_format()
            cell_format.set_font_name('Verdana')
            cell_format.set_font_size(11)
            worksheet.set_column("A:H", None, cell_format)
            # Columns
            worksheet.set_column("A:A", 4.5)
            worksheet.set_column("B:B", 19)
            worksheet.set_column("C:F", 7)
            worksheet.set_column("G:G", 9)
            worksheet.set_column("H:H", 15)
        elif option == 4:
            # format-font
            cell_format = workbook.add_format()
            cell_format.set_font_name('Verdana')
            cell_format.set_font_size(11)
            worksheet.set_column("A:J", None, cell_format)
            # Columns
            worksheet.set_column("A:A", 4.5)
            worksheet.set_column("B:B", 19)
            worksheet.set_column("C:D", 7)
            worksheet.set_column("E:E", 11.43)            
            worksheet.set_column("F:G", 7)
            worksheet.set_column("H:H", 17)
            worksheet.set_column("I:I", 21)
            worksheet.set_column("J:J", 19)
        elif option == 5:
            # format-font
            cell_format = workbook.add_format()
            cell_format.set_font_name('Verdana')
            cell_format.set_font_size(11)
            worksheet.set_column("A:K", None, cell_format)
            # Columns
            worksheet.set_column("A:A", 4.5)
            worksheet.set_column("B:B", 19)
            worksheet.set_column("C:D", 7)
            worksheet.set_column("E:E", 11.43)            
            worksheet.set_column("F:G", 7)
            worksheet.set_column("H:H", 17)
            worksheet.set_column("I:I", 57)
            worksheet.set_column("J:J", 19)
            worksheet.set_column("K:K", 17)
        elif option == 6:
            # format-font
            cell_format = workbook.add_format()
            cell_format.set_font_name('Verdana')
            cell_format.set_font_size(11)
            worksheet.set_column("A:L", None, cell_format)
            # Columns
            worksheet.set_column("A:A", 4.5)
            worksheet.set_column("B:B", 19)
            worksheet.set_column("C:D", 7)
            worksheet.set_column("E:E", 11.43)            
            worksheet.set_column("F:G", 10)
            worksheet.set_column("H:H", 5.71)
            worksheet.set_column("I:I", 8.29)
            worksheet.set_column("J:J", 19)
            worksheet.set_column("K:K", 17)
            worksheet.set_column("L:L", 16)           
            
                       
        (max_row, max_col) = df.shape
        # Create a list of column headers, to use in add_table().
        column_settings = [{'header': column} for column in df.columns]
        # Add the Excel table structure. Pandas will add the data.
        worksheet.add_table(0, 1, max_row, max_col, {'columns': column_settings})
        # Make the columns wider for clarity.
        #worksheet.set_column(0, max_col + 1, 4.5)  

        writer.save()
              
    else:
        if option == 1 or option == 2 or option == 6: 
            print(f"\nEl buque seleccionado {seleccion} no coincide con la búsqueda o no tiene unidades vacías para asignar una secuencia de retiro ")
            ornament()
        elif option == 3:
            print(f"\nLa línea seleccionada {seleccion} no se encuentra en el archivo")
            ornament()
        elif option == 4:
            print(f"\nEl puerto seleccionado {seleccion} no se encuentra en el archivo")
            ornament()
        elif option == 5:
           print(f"\nLa retención seleccionada {seleccion} no se encuentra en el archivo")
           ornament()    
        

def ornament():
    print("_" * espacios)


def menu1():
    ornament()
    print("Planning Yard")
    ornament()
      
    
def menu2():
    #print("\n")
    ornament()
    print("Opciones disponibles:")
    ornament()
    print("1. Listado vacíos descargados desde Motonaves")
    print("2. Listado vacíos desde el FreePool")
    print("3. Listado vacíos por Línea")
    print("4. Listado vacíos por Puerto")
    print("5. Listado vacíos por Retención")
    print ("6. Asignación Evacuación")
    print ("7. Digitar Posicionados")
    print ("8. Ocupación")
    print("0. Salir")
    print()


def menu3():
    ornament()
    print("Adicionar más Puertos?:")
    print()
    print("1. Continuar sin adicionar otro Puerto")
    print("2. Agregar otro Puerto de Descarga:")
    ornament()
    print()

  
def main():  # MAIN
    root = Tk() # ERASE DIALOG
    root.update()
    clear()  
    menu1()
    seg = 2.5
    excel_file= openFile()  # Name of the file to choose
    root.destroy() # Close Dialog
    wb = xlrd.open_workbook(excel_file, encoding_override="cp1252") # Read excel book
    sheet = wb.sheet_by_index(0)

    while True:       
            option = menu2()
            try:
                option =  int(input("Seleccione una opción: "))
                ornament()
                if option == 1: # Vacíos desde buque
                    seleccion =  askUser(option)
                    ordered_list(sheet, seleccion, option)
                    #break
                elif option == 2: # Vacíos desde freepool
                    seleccion =  askUser(option)
                    ordered_list(sheet, seleccion, option)
                    #break
                elif option == 3:
                    seleccion =  askUser(option)
                    ordered_list(sheet, seleccion, option)
                    #break
                elif option == 4:
                    seleccion =  askUser(option)
                    ordered_list(sheet, seleccion, option)
                    #break
                elif option == 5: # Retención
                    seleccion =  askUser(option)
                    ordered_list(sheet, seleccion, option)
                    time.sleep(seg)
                    clear()
                elif option == 6: # Asignación Vacíos
                    seleccion =  askUser(option)
                    ordered_list(sheet, seleccion, option)
                    time.sleep(seg)
                elif option == 7: #Posicionados
                    root = Tk() # ERASE DIALOG
                    root.update()
                    pos_list = getFile(pathFile)
                    root.destroy() 
                    typeContainer(pos_list)
                    time.sleep(seg)
                    clear()
                elif option == 8: # Ocupación
                    ornament()
                    print("OCUPACIÓN EN TEUS".center(espacios))
                    ornament()
                    ocupacion = ocupacion_list(sheet, pathFile)                              
                elif option == 0:
                    break
                else:
                    print()
                    clear()
                    print("Error, solo de aceptan números del 0 al 8")
            except ValueError:
            #except TypeError as err:       
                  print("Error, ingrese solamente números")#, err)            
    
    


if __name__ == '__main__':
    main()


os.system("pause") # Press a key to continue 
